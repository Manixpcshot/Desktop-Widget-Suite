"""Create Persian, RTL release notes from the latest 20 Git commits."""
import re
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def generate(version: str) -> Path:
    if not re.fullmatch(r"v\d+\.\d+\.\d+(?:-[A-Za-z0-9.-]+)?", version):
        raise ValueError("Expected a version such as v0.1.0")
    commits = subprocess.check_output(
        ["git", "log", "-20", "--format=%h%x09%s"], text=True, encoding="utf-8"
    ).splitlines()
    wb = Workbook()
    ws = wb.active
    ws.title = "Release Notes"
    ws.sheet_view.rightToLeft = True
    ws.append(["نسخه", "تاریخ انتشار (UTC)", "شماره Commit", "توضیحات تغییر"])
    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")
    for commit in commits:
        sha, _, subject = commit.partition("\t")
        ws.append([version, timestamp, sha, subject])
        # Git commit subjects are untrusted text, never spreadsheet formulas.
        for cell in ws[ws.max_row]:
            cell.data_type = "s"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="253047")
    for column, width in {"A": 20, "B": 25, "C": 18, "D": 100}.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    path = Path(f"release_notes_{version}.xlsx")
    wb.save(path)
    return path


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: generate_release_excel.py vX.Y.Z")
    print(generate(sys.argv[1]))
